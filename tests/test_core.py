from __future__ import annotations

import io
import unittest
from pathlib import Path

from openpyxl import load_workbook
from reportlab.pdfgen import canvas

from approval_app.core import PdfValidationError, compare_documents, extract_document
from approval_app.exporters import export_excel, export_pdf

ROOT = Path(__file__).resolve().parents[1]


def make_pdf(lines: list[str]) -> bytes:
    output = io.BytesIO()
    pdf = canvas.Canvas(output)
    y = 800
    for line in lines:
        pdf.drawString(50, y, line)
        y -= 18
    pdf.save()
    return output.getvalue()


class SampleIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.old = extract_document((ROOT / "IEC1000-07N.pdf").read_bytes(), "IEC1000-07N.pdf")
        cls.new = extract_document((ROOT / "IEC1000-07NB1.pdf").read_bytes(), "IEC1000-07NB1.pdf")
        cls.result = compare_documents(cls.old, cls.new)

    def test_expected_structured_changes(self):
        changed = {item.item: (item.old_value, item.new_value) for item in self.result.changes}
        self.assertEqual(changed["모델번호"], ("IEC1000-07N", "IEC1000-07NB1"))
        self.assertEqual(changed["등록번호"], ("HNS-10007N", "HNS-10007NB1"))
        self.assertEqual(changed["제품무게"], ("540g", "630g"))

    def test_expected_visual_pages_only(self):
        visual_pages = {item.new_page for item in self.result.changes if item.change_type == "시각 변경"}
        self.assertEqual(visual_pages, {3, 4, 5})
        status = {item.new_page: item.status for item in self.result.pages}
        self.assertEqual(status[2], "동일")
        self.assertEqual(status[7], "동일")
        self.assertEqual(status[8], "동일")
        self.assertEqual(status[9], "동일")

    def test_expected_quality_guides(self):
        keys = {(issue.item, issue.issue_type) for issue in self.result.issues}
        self.assertIn(("Optimun", "오타"), keys)
        self.assertIn(("Moblie", "오타"), keys)
        self.assertIn(("IIC / I2C", "표기 불일치"), keys)
        self.assertIn(("모델 참조", "모델·등록번호 불일치"), keys)
        model_issue = next(issue for issue in self.result.issues if issue.item == "모델 참조")
        self.assertEqual(model_issue.page, 6)

    def test_negative_temperature_is_preserved(self):
        fields = {(field.key, field.page): field.value for field in self.new.fields}
        self.assertEqual(fields[("보관온도", 6)], "-30 ~ 70°C")

    def test_exports_are_valid(self):
        rows = [row.to_dict() for row in self.result.review_rows]
        excel = export_excel(self.result, rows)
        pdf = export_pdf(self.result, rows)
        workbook = load_workbook(io.BytesIO(excel), read_only=True)
        self.assertEqual(workbook.sheetnames, ["변경 및 검수", "페이지 비교", "처리 정보"])
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertGreater(len(pdf), 1000)


class ValidationAndDiffTests(unittest.TestCase):
    def test_rejects_non_pdf_extension(self):
        with self.assertRaises(PdfValidationError):
            extract_document(b"not a pdf", "sample.txt")

    def test_rejects_empty_pdf(self):
        with self.assertRaises(PdfValidationError):
            extract_document(b"", "sample.pdf")

    def test_page_added(self):
        old = extract_document(make_pdf(["Model No IEC1000-01N", "registration"]), "old.pdf")
        new_data = io.BytesIO()
        pdf = canvas.Canvas(new_data)
        pdf.drawString(50, 800, "Model No IEC1000-01N")
        pdf.showPage()
        pdf.drawString(50, 800, "added page")
        pdf.save()
        new = extract_document(new_data.getvalue(), "new.pdf")
        result = compare_documents(old, new)
        self.assertTrue(any(change.change_type == "추가" and change.item == "2페이지" for change in result.changes))


if __name__ == "__main__":
    unittest.main()
